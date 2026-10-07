"""增量报送（watch）、内容分类与 DOCX 日报的离线测试。

    python -m unittest discover -s tests -v
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from frisbee_radar.categories import Category, CategoryClassifier  # noqa: E402
from frisbee_radar.config import (  # noqa: E402
    CategoryMatchSettings,
    OutputSettings,
    WatchSettings,
    load_config,
)
from frisbee_radar.docx_report import build_docx_daily  # noqa: E402
from frisbee_radar.models import Post  # noqa: E402
from frisbee_radar.storage import Storage  # noqa: E402

AUTH = "全国飞盘运动推广委员会"

# 测试用的迷你词表，形状和 config/sources.yaml 里的一致
TRAINING = Category(
    key="training", name="训练与技巧", order=1,
    include=["教学", "技巧", "训练", "入门", "反手", "传盘", "战术", "规则"],
    exclude=["俱乐部", "战队", "联赛", "高校", "大学"],
)
TEAM = Category(
    key="team", name="队伍与赛事动态", order=2,
    include=["俱乐部", "战队", "招新", "迎新", "联赛", "省赛", "交流赛", "校队", "高校"],
    exclude=["为什么", "怎么样", "凉了", "火了", "出圈", "争议", "媛"],
)


def make_post(post_id: str, **overrides) -> Post:
    base = {
        "platform": "wechat",
        "post_id": post_id,
        "title": f"飞盘俱乐部活动 {post_id}",
        "content": "本周六下午在滨江公园有一场新手友好的极限飞盘活动，欢迎参加。",
        "author": "某飞盘俱乐部",
        "url": f"https://mp.weixin.qq.com/s/{post_id}",
        "publish_time": datetime.now(timezone.utc),
        "relevance": 0.8,
    }
    base.update(overrides)
    return Post(**base)


# ============================================================ 增量报送


class TestIncrementalTracking(unittest.TestCase):
    """reported_ts 是 watch 判断「有没有新东西」的唯一依据。"""

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.storage = Storage(Path(self.tmp.name) / "inc.db")

    def tearDown(self) -> None:
        self.storage.close()
        self.tmp.cleanup()

    def test_everything_starts_unreported(self):
        self.storage.upsert_many([make_post("a"), make_post("b")])
        self.assertEqual(self.storage.count_unreported(0.5), 2)
        self.assertEqual(len(self.storage.unreported(0.5)), 2)

    def test_marking_removes_from_pending(self):
        self.storage.upsert_many([make_post("a"), make_post("b")])
        pending = self.storage.unreported(0.5)
        self.storage.mark_reported(pending[:1])
        self.assertEqual(self.storage.count_unreported(0.5), 1)
        self.assertEqual([p.post_id for p in self.storage.unreported(0.5)], ["b"])

    def test_marking_is_idempotent(self):
        self.storage.upsert_many([make_post("a")])
        posts = self.storage.unreported(0.5)
        self.storage.mark_reported(posts)
        self.storage.mark_reported(posts)
        self.assertEqual(self.storage.count_unreported(0.5), 0)

    def test_irrelevant_content_never_becomes_pending(self):
        self.storage.upsert_many([
            make_post("good", relevance=0.8),
            make_post("noise", relevance=0.0),
        ])
        self.assertEqual(self.storage.count_unreported(0.5), 1)

    def test_authoritative_filtering(self):
        self.storage.upsert_many([
            make_post("x1", author=AUTH),
            make_post("x2", author=AUTH),
            make_post("y1", author="路人号"),
        ])
        self.assertEqual(self.storage.count_unreported(0.5, authors=[AUTH]), 2)
        self.assertEqual(len(self.storage.unreported_from_authors([AUTH], 0.5)), 2)
        # 没配权威号时不应误伤
        self.assertEqual(self.storage.unreported_from_authors([], 0.5), [])
        self.assertEqual(self.storage.count_unreported(0.5, authors=[]), 3)

    def test_last_reported_at_is_none_before_first_report(self):
        self.storage.upsert_many([make_post("a")])
        self.assertIsNone(self.storage.last_reported_at())

    def test_last_reported_at_tracks_latest_report(self):
        self.storage.upsert_many([make_post("a")])
        self.storage.mark_reported(self.storage.unreported(0.5))
        moment = self.storage.last_reported_at()
        self.assertIsNotNone(moment)
        self.assertLess(abs((datetime.now(timezone.utc) - moment).total_seconds()), 60)

    def test_updating_an_existing_post_does_not_resurrect_it(self):
        """重跑采集会刷新互动数据，但不该让已报送的内容重新排队。"""
        self.storage.upsert_many([make_post("a")])
        self.storage.mark_reported(self.storage.unreported(0.5))
        self.assertEqual(self.storage.count_unreported(0.5), 0)

        self.storage.upsert_many([make_post("a", likes=999)])
        self.assertEqual(self.storage.count_unreported(0.5), 0)
        self.assertEqual(self.storage.query()[0].likes, 999)

    def test_baseline_clears_backlog(self):
        self.storage.upsert_many([make_post(f"p{i}") for i in range(5)])
        self.assertEqual(self.storage.count_unreported(0.5), 5)
        cleared = self.storage.mark_all_reported_before()
        self.assertEqual(cleared, 5)
        self.assertEqual(self.storage.count_unreported(0.5), 0)
        # 基线只影响报送标记，数据一条都不少
        self.assertEqual(self.storage.stats()["_total"], 5)

    def test_author_counts(self):
        self.storage.upsert_many([
            make_post("a1", author="甲"),
            make_post("a2", author="甲"),
            make_post("b1", author="乙"),
            make_post("c1", author=""),
        ])
        counts = dict(self.storage.author_counts())
        self.assertEqual(counts.get("甲"), 2)
        self.assertEqual(counts.get("乙"), 1)
        self.assertNotIn("", counts)   # 空作者不参与统计

    def test_migration_adds_missing_columns(self):
        """老库（没有 reported_ts）打开时应当自动补列，而不是报错。"""
        db = Path(self.tmp.name) / "legacy.db"
        import sqlite3

        conn = sqlite3.connect(str(db))
        conn.executescript(
            """CREATE TABLE posts (
                   platform TEXT NOT NULL, post_id TEXT NOT NULL,
                   title TEXT DEFAULT '', content TEXT DEFAULT '',
                   author TEXT DEFAULT '', author_id TEXT DEFAULT '',
                   url TEXT DEFAULT '', publish_ts INTEGER,
                   likes INTEGER DEFAULT 0, comments INTEGER DEFAULT 0,
                   shares INTEGER DEFAULT 0, collects INTEGER DEFAULT 0,
                   views INTEGER DEFAULT 0, images TEXT DEFAULT '[]',
                   videos TEXT DEFAULT '[]', collected_ts INTEGER,
                   matched_by TEXT DEFAULT '', relevance REAL DEFAULT 0.0,
                   relevance_hits TEXT DEFAULT '[]', content_hash TEXT DEFAULT '',
                   raw TEXT DEFAULT '{}',
                   PRIMARY KEY (platform, post_id));"""
        )
        conn.execute(
            "INSERT INTO posts (platform, post_id, title, relevance) VALUES ('wechat','old','旧内容',0.9)"
        )
        conn.commit()
        conn.close()

        with Storage(db) as migrated:
            columns = {r["name"] for r in migrated.conn.execute("PRAGMA table_info(posts)")}
            self.assertIn("reported_ts", columns)
            # 老数据默认未报送，会进第一份日报——所以要先 baseline
            self.assertEqual(migrated.count_unreported(0.5), 1)


# ============================================================ 配置


class TestWatchConfig(unittest.TestCase):
    def test_defaults_are_conservative(self):
        cfg = WatchSettings.from_dict(None)
        self.assertEqual(cfg.authoritative, [])
        self.assertGreater(cfg.trigger_min_new, 0)
        self.assertGreater(cfg.min_interval_hours, 0)

    def test_parsing(self):
        cfg = WatchSettings.from_dict({
            "authoritative": ["  某号  ", "", AUTH],
            "trigger_min_new": 3,
            "min_interval_hours": 1.5,
            "authoritative_breaks_interval": False,
        })
        self.assertEqual(cfg.authoritative, ["某号", AUTH])   # 去空白、丢空串
        self.assertEqual(cfg.trigger_min_new, 3)
        self.assertFalse(cfg.authoritative_breaks_interval)

    def test_shipped_config_loads(self):
        cfg = load_config()
        self.assertTrue(cfg.watch.authoritative, "sources.yaml 应当已配权威号")
        self.assertGreater(cfg.output.max_items, 0)
        self.assertTrue(cfg.output.title)
        self.assertTrue(cfg.categories, "sources.yaml 应当配了分类")

    def test_shipped_categories_are_ordered_training_first(self):
        cfg = load_config()
        ordered = sorted(cfg.categories, key=lambda c: c.order)
        self.assertEqual(ordered[0].key, "training")

    def test_output_filename_and_dir(self):
        cfg = OutputSettings()
        moment = datetime(2026, 10, 5)
        self.assertEqual(cfg.resolve_filename(moment), "2026-10-05_飞盘日报.docx")
        self.assertTrue(cfg.resolve_dir().is_absolute())

        # 绝对路径原样返回，相对路径按项目根解析
        absolute = Path("C:/abs/path")
        self.assertEqual(OutputSettings(dir=str(absolute)).resolve_dir(), absolute)
        self.assertTrue(OutputSettings(dir="data/reports").resolve_dir().is_absolute())

    def test_category_match_defaults_to_title_only(self):
        """正文兜底默认关闭——开了会把媒体稿误收进来。"""
        self.assertEqual(CategoryMatchSettings.from_dict(None).body_min_hits, 0)
        self.assertEqual(load_config().category_match.body_min_hits, 0)


# ============================================================ 分类


class TestCategoryClassifier(unittest.TestCase):
    """分类是白名单：命中不了任何分类的内容不进日报。"""

    def setUp(self) -> None:
        self.clf = CategoryClassifier([TRAINING, TEAM])

    def test_training_title_wins_by_order(self):
        match = self.clf.classify("飞盘反手传盘教学")
        self.assertIsNotNone(match)
        self.assertEqual(match.category.key, "training")
        self.assertEqual(match.matched_in, "title")

    def test_team_title_classified(self):
        match = self.clf.classify("我校飞盘俱乐部招新啦")
        self.assertIsNotNone(match)
        self.assertEqual(match.category.key, "team")

    def test_uncategorized_is_dropped(self):
        # 媒体评论类，两类都不该收
        for title in ("为什么年轻人都不玩飞盘了?", "飞盘火了,VC开始投瑜伽裤"):
            self.assertIsNone(self.clf.classify(title), title)

    def test_exclude_beats_include(self):
        """「掷准飞盘俱乐部」不能被「训练」类拉走，应归队伍类。"""
        clf = CategoryClassifier([
            Category(key="training", name="训练", order=1,
                     include=["掷准"], exclude=["俱乐部"]),
            Category(key="team", name="队伍", order=2, include=["俱乐部"]),
        ])
        match = clf.classify("一个神秘掷准飞盘俱乐部现身")
        self.assertIsNotNone(match)
        self.assertEqual(match.category.key, "team")

    def test_title_only_by_default(self):
        """标题没信号时，正文提到训练也不该收——精度优先。"""
        self.assertIsNone(self.clf.classify("周末一起运动", "我们有教练带训练"))

    def test_body_fallback_when_enabled(self):
        clf = CategoryClassifier([TRAINING, TEAM], body_min_hits=2)
        match = clf.classify("周末一起运动", "我们有教练带队训练，还会教传盘")
        self.assertIsNotNone(match)
        self.assertEqual(match.matched_in, "content")

    def test_global_exclude_blocks_everything(self):
        clf = CategoryClassifier([TRAINING], global_exclude=["硬盘"])
        self.assertIsNone(clf.classify("飞盘教学：硬盘里的飞盘"))
        self.assertIsNotNone(clf.classify("飞盘教学：反手"))

    def test_whitespace_insensitive(self):
        """「飞 盘 教 学」这种空格强调也要能命中。"""
        self.assertIsNotNone(self.clf.classify("飞 盘 教 学"))

    def test_split_returns_excluded_separately(self):
        posts = [
            make_post("t1", title="飞盘反手传盘教学", content=""),
            make_post("m1", title="为什么年轻人都不玩飞盘了?", content=""),
        ]
        grouped, excluded = self.clf.split(posts)
        self.assertEqual(len(grouped["training"]), 1)
        self.assertEqual(len(excluded), 1)

    def test_group_sorts_newest_first(self):
        now = datetime.now(timezone.utc)
        posts = [
            make_post("old", title="飞盘教学 A", content="",
                      publish_time=now - timedelta(days=100)),
            make_post("new", title="飞盘教学 B", content="", publish_time=now),
            make_post("mid", title="飞盘教学 C", content="",
                      publish_time=now - timedelta(days=10)),
        ]
        grouped = self.clf.group(posts)
        self.assertEqual([p.post_id for p in grouped["training"]], ["new", "mid", "old"])

    def test_group_keeps_undated_items_at_the_end(self):
        now = datetime.now(timezone.utc)
        posts = [
            make_post("nodate", title="飞盘教学 A", content="", publish_time=None),
            make_post("dated", title="飞盘教学 B", content="", publish_time=now),
        ]
        grouped = self.clf.group(posts)
        self.assertEqual([p.post_id for p in grouped["training"]], ["dated", "nodate"])

    def test_empty_include_category_is_ignored(self):
        clf = CategoryClassifier([Category(key="empty", name="空", order=1, include=[])])
        self.assertIsNone(clf.classify("飞盘教学"))


# ============================================================ DOCX 日报


class TestDocxDaily(unittest.TestCase):
    """生成 DOCX 并校验结构。视觉判断靠渲染出图，不在这层。"""

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.now = datetime.now(timezone.utc)
        self.posts = [
            make_post("t-early", title="飞盘基础教学课程", content="教练讲解反手传盘要领。",
                      relevance=0.9, publish_time=self.now - timedelta(days=90)),
            make_post("t-late", title="飞盘战术配合训练", content="阵型与跑位练习。",
                      relevance=0.9, publish_time=self.now - timedelta(days=1)),
            make_post("team-1", title="我校飞盘俱乐部招新啦", content="招新信息与联系方式。",
                      author=AUTH, relevance=0.85,
                      publish_time=self.now - timedelta(days=3)),
            make_post("media-1", title="为什么年轻人都不玩飞盘了?", content="行业观察。",
                      relevance=0.7, publish_time=self.now),
        ]

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _build(self, posts=None, **kwargs):
        path = Path(self.tmp.name) / "daily.docx"
        build_docx_daily(
            posts if posts is not None else self.posts,
            path,
            classifier=CategoryClassifier([TRAINING, TEAM]),
            **kwargs,
        )
        return path

    def _text(self, path) -> str:
        from docx import Document

        doc = Document(str(path))
        chunks = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    chunks.append(cell.text)
        return "\n".join(chunks)

    def test_creates_readable_docx(self):
        path = self._build()
        self.assertTrue(path.exists())
        self.assertGreater(path.stat().st_size, 3000)

    def test_single_section_no_cover_no_toc(self):
        """演示幕是紧凑题头，不是整页封面；也不需要目录页。"""
        from docx import Document

        doc = Document(str(self._build()))
        self.assertEqual(len(doc.sections), 1)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("目　录", text)
        self.assertNotIn("目录", text)

    def test_masthead_block_present(self):
        from docx import Document

        doc = Document(str(self._build()))
        masthead = doc.tables[0]
        self.assertEqual(len(masthead.rows), 1)
        self.assertIn("飞盘日报", masthead.rows[0].cells[0].text)

    def test_training_section_before_team_section(self):
        """用户要求：训练技巧优先，队伍日常次之。"""
        text = self._text(self._build())
        self.assertLess(text.index("训练与技巧"), text.index("队伍与赛事动态"))

    def test_items_sorted_newest_first_within_section(self):
        text = self._text(self._build())
        self.assertLess(
            text.index("飞盘战术配合训练"),      # 1 天前
            text.index("飞盘基础教学课程"),      # 90 天前
        )

    def test_uncategorized_content_excluded(self):
        text = self._text(self._build())
        self.assertNotIn("为什么年轻人都不玩飞盘了?", text)

    def test_authoritative_marked_in_meta(self):
        text = self._text(self._build(authoritative=[AUTH]))
        self.assertIn("权威号", text)

    def test_authoritative_marker_absent_when_not_configured(self):
        """没配权威号时不该凭空出现这个标记。"""
        text = self._text(self._build(authoritative=[]))
        self.assertNotIn("权威号", text)

    def test_headings_use_real_heading_styles(self):
        from docx import Document

        doc = Document(str(self._build()))
        styles = [p.style.name for p in doc.paragraphs if p.style.name.startswith("Heading")]
        self.assertIn("Heading 1", styles)

    def test_headings_have_outline_level(self):
        from docx import Document
        from docx.oxml.ns import qn

        doc = Document(str(self._build()))
        pPr = doc.styles["Heading 1"].element.find(qn("w:pPr"))
        self.assertIsNotNone(pPr)
        self.assertIsNotNone(pPr.find(qn("w:outlineLvl")))

    def test_footer_page_number_uses_word_switch(self):
        """footer 写 `\\* decimal` 会渲染成「1decimal」，必须用 arabic。"""
        from docx import Document

        doc = Document(str(self._build()))
        xml = doc.sections[0].footer._element.xml
        self.assertIn("arabic", xml)
        self.assertNotIn("decimal", xml)

    def test_table_widths_written_to_grid(self):
        """只设 cell.width 不够——缺 tblGrid 列宽会被当成等分。"""
        from docx import Document
        from docx.oxml.ns import qn

        doc = Document(str(self._build()))
        table = doc.tables[1]      # 0=题头 1=概览表
        grid = table._tbl.find(qn("w:tblGrid"))
        self.assertIsNotNone(grid, "表格缺 tblGrid，列宽不会生效")

        # 概览表列宽是 [4.5, 2.0, 3.5, 3.5] cm；后两列等宽是对的，
        # 所以这里校验换算后的绝对宽度，而不是「互不相同」
        cols = [int(gc.get(qn("w:w"))) for gc in grid.findall(qn("w:gridCol"))]
        self.assertEqual(cols, [int(w * 567) for w in (4.5, 2.0, 3.5, 3.5)])
        self.assertNotEqual(len(set(cols)), 1, "列宽全相同说明没写进去，被当等分了")
        self.assertLess(int(sum(cols) / 567), 15.6, "表格总宽不该超出正文宽度")

        layout = table._tbl.tblPr.find(qn("w:tblLayout"))
        self.assertEqual(layout.get(qn("w:type")), "fixed")

    def test_links_use_short_label_not_raw_url(self):
        """公众号链接 200+ 字符，直接铺出来会把一行撑满。"""
        text = self._text(self._build())
        self.assertIn("打开原文", text)
        self.assertNotIn("mp.weixin.qq.com/s?src=", text)

    def test_shading_uses_clear_not_solid(self):
        import zipfile

        with zipfile.ZipFile(self._build()) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        self.assertNotIn('w:val="solid"', xml)

    def test_no_markdown_artifacts(self):
        body = self._text(self._build())
        for artifact in ("**", "|---", "## ", "```"):
            self.assertNotIn(artifact, body, f"文档里出现 Markdown 残留：{artifact!r}")

    def test_no_text_character_rule_lines(self):
        body = self._text(self._build())
        for line in ("───", "━━━", "═══", "——————"):
            self.assertNotIn(line, body)

    def test_masthead_block_has_no_borders(self):
        """默认表格边框会给有色底块加高，必须显式去掉。"""
        from docx import Document
        from docx.oxml.ns import qn

        doc = Document(str(self._build()))
        borders = doc.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
        self.assertIsNotNone(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            self.assertEqual(borders.find(qn(f"w:{edge}")).get(qn("w:val")), "none")

    def test_empty_input_produces_valid_docx_with_explanation(self):
        from docx import Document

        doc = Document(str(self._build(posts=[])))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("没有可收录的内容", text)
        self.assertEqual(len(doc.sections), 1)

    def test_data_notes_report_scanned_and_excluded_counts(self):
        """数据说明要写清「扫了多少、剔了多少」，不能是同义反复。"""
        text = self._text(self._build(scanned_total=10, excluded_total=6))
        self.assertIn("扫描到 10 条", text)
        self.assertIn("另有 6 条", text)

    def test_body_paragraphs_indented_and_justified(self):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        doc = Document(str(self._build()))
        indented = 0
        for para in doc.paragraphs:
            if para.style.name != "Normal" or len(para.text) < 40:
                continue
            if para.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
                pPr = para._p.find(qn("w:pPr"))
                if pPr is not None and pPr.find(qn("w:ind")) is not None:
                    indented += 1
        self.assertGreater(indented, 0, "正文段落应首行缩进 + 两端对齐")


if __name__ == "__main__":
    unittest.main(verbosity=2)
