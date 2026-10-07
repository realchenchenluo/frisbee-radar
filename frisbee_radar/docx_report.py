"""DOCX 日报生成。

排版口径参考用户已有的「X 情报日报」：**不做封面、不做目录**，第一页
顶部一个紧凑题头，然后直接进正文往下读。对一份每天都在看的简报来说，
整页封面和目录页都是每次阅读都要翻过去的负担。

字体仍按 office skill 的正式体规范：
    H1 黑体 16pt / H2 黑体 14pt / 正文宋体 12pt / 说明性文字 10.5pt
    行距 1.3 倍、正文首行缩进 2 字符、纯黑字色
    A4，页边距 上下 2.54cm / 左 3.0cm / 右 2.5cm

内容顺序（用户明确要求）：
    一、训练与技巧        组内按发布时间倒序
    二、队伍与赛事动态     同上
    不属于任何分类的内容不进日报 —— 见 categories.py

⚠️ 两条踩过坑的约束，改这个文件时别踩：
  1. 表格必须写 w:tblGrid（只设 cell.width 无效，列宽会被当等分）。
  2. 不要用字符画的分割线（───），跨渲染器宽度不一致，用段落边框。
"""

from __future__ import annotations

from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .categories import CategoryClassifier
from .models import Post, platform_label

# ---------------------------------------------------------------- 设计令牌


class Palette:
    """FG-1 Forest Mint —— 户外运动语境。"""

    brand_bg = "0C1F1A"        # 题头区块底色
    brand_accent = "3DDBB5"    # 强调线 / 标签
    brand_text = "FFFFFF"
    brand_sub = "B0B8C0"

    heading_text = "000000"
    body_text = "000000"
    meta_text = "5A6570"       # 元信息（来源/日期/互动）
    summary_text = "404850"    # 摘要
    link_text = "1B6B7A"

    table_header_bg = "2A7A65"
    table_header_text = "FFFFFF"
    rule_light = "D8E0DC"      # 条目之间的细分割线


CN_HEI = "SimHei"
CN_SONG = "SimSun"
CN_YAHEI = "Microsoft YaHei"
EN_SERIF = "Times New Roman"
EN_SANS = "Arial"

SIZE_DOC_TITLE = 22
SIZE_H1, SIZE_H2 = 16, 14
SIZE_BODY = 12
SIZE_ITEM_TITLE = 11.5
SIZE_META = 9
SIZE_SUMMARY = 10
SIZE_CAPTION = 10.5

LINE_SPACING = 1.3
BODY_INDENT = Pt(24)       # 480 缇 = 2 字符 @宋体12pt

PAGE_W, PAGE_H = Cm(21), Cm(29.7)
MARGIN_TB, MARGIN_L, MARGIN_R = Cm(2.54), Cm(3.0), Cm(2.5)
TEXT_WIDTH_CM = 21.0 - 3.0 - 2.5      # 15.5cm，表格总宽按它算


def _hex(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#").upper())


def _fmt_count(value: int) -> str:
    if value >= 100_000_000:
        return f"{value / 100_000_000:.1f}亿"
    if value >= 10_000:
        return f"{value / 10_000:.1f}万"
    return f"{value:,}"


def _fmt_date(dt: datetime | None) -> str:
    if dt is None:
        return "时间未知"
    return dt.astimezone(timezone(timedelta(hours=8))).strftime("%Y-%m-%d")


def _clean(text: str, limit: int = 0) -> str:
    """压平空白。同时兜住 None/空值，避免文档里出现字面 undefined/null。"""
    flat = " ".join(str(text or "").split())
    if limit and len(flat) > limit:
        return flat[: limit - 1] + "…"
    return flat


# ---------------------------------------------------------------- 底层构件


def _set_run(run, *, cn=CN_SONG, en=EN_SERIF, size=SIZE_BODY,
             bold=False, color=Palette.body_text, italic=False) -> None:
    """设置西文 + 中文字体。

    python-docx 的 run.font.name 只写 w:ascii/w:hAnsi，中文会走默认字体，
    必须单独设 w:eastAsia 才生效。
    """
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _hex(color)

    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"), en)


def _set_spacing(paragraph, *, before=0, after=0, line=None,
                 rule=WD_LINE_SPACING.MULTIPLE) -> None:
    pf = paragraph.paragraph_format
    if before:
        pf.space_before = Pt(before / 20)
    if after:
        pf.space_after = Pt(after / 20)
    if line is None:
        pf.line_spacing = LINE_SPACING          # → w:line=312 lineRule=auto
    elif rule is WD_LINE_SPACING.MULTIPLE:
        pf.line_spacing = line
    else:
        pf.line_spacing_rule = rule
        pf.line_spacing = Pt(line / 20)


def _para_border(paragraph, *, edge: str, size=6, color=Palette.brand_accent,
                 space=8) -> None:
    """段落边框当装饰线。规范禁止 ─── 这类字符画线。"""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        shd = pPr.find(qn("w:shd"))
        if shd is not None:
            shd.addprevious(pBdr)
        else:
            pPr.append(pBdr)

    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color.lstrip("#").upper())
    pBdr.append(el)


def _shade(element, fill: str) -> None:
    """单元格底色。必须 clear —— solid 在 WPS 里会渲染成纯黑。"""
    tcPr = element.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill.lstrip("#").upper())


def _no_borders(table) -> None:
    """去掉表格全部边框。有色底的区块表必须显式去边，否则默认边框会加高。"""
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    table._tbl.tblPr.append(borders)

    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            cb = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:val"), "none")
                el.set(qn("w:sz"), "0")
                cb.append(el)
            tcPr.append(cb)


def _set_table_widths(table, widths_cm: list[float]) -> None:
    """固定列宽。

    只设 cell.width 不够：Word/WPS/LibreOffice 都优先读 w:tblGrid，
    缺了它列宽会被当等分（标题列会被挤成三行）。所以 gr id、tblW、
    tblLayout(fixed)、缩进都要一起写。
    """
    tbl = table._tbl
    tblPr = tbl.tblPr

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)

    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    tblPr.append(ind)

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(width * 567)))
        grid.append(gc)
    tblPr.addnext(grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths_cm):
                cell.width = Cm(widths_cm[idx])


def _cell_margins(table, top=60, bottom=60, left=110, right=110) -> None:
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


def _row_cant_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _repeat_header(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def _keep_with_next(paragraph) -> None:
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))


def _add_hyperlink(paragraph, url: str, text: str, size=SIZE_META) -> None:
    """插入真实可点击的超链接。

    公众号链接长达 200+ 字符，直接铺出来会把一行撑满还看不清，
    所以只在 href 里放完整地址，显示文字用短标签。
    """
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi"):
        rFonts.set(qn(attr), EN_SANS)
    rFonts.set(qn("w:eastAsia"), CN_SONG)
    rPr.append(rFonts)
    for tag, val in (("w:color", Palette.link_text.lstrip("#").upper()),):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        rPr.append(el)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    run.append(rPr)

    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _footer_page_number(section) -> None:
    """页脚页码。

    域开关必须写 Word 的写法（arabic），不能写 docx-js 的枚举名 decimal
    —— 写 decimal 会渲染成「1decimal」。
    """
    section.footer.is_linked_to_previous = False
    footer = section.footer
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in list(para.runs):
        run._r.getparent().remove(run._r)

    run = para.add_run()
    _set_run(run, size=SIZE_META, color=Palette.meta_text)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r"PAGE \* arabic \* MERGEFORMAT"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, end):
        run._r.append(el)


def _set_outline_level(style, level: int) -> None:
    pPr = style.element.get_or_add_pPr()
    ol = pPr.find(qn("w:outlineLvl"))
    if ol is None:
        ol = OxmlElement("w:outlineLvl")
        pPr.append(ol)
    ol.set(qn("w:val"), str(level))


# ---------------------------------------------------------------- 日报生成


class DocxDailyReport:
    """把一批 Post 渲染成一份日报。"""

    def __init__(
        self,
        posts: list[Post],
        *,
        title: str = "飞盘日报",
        classifier: CategoryClassifier | None = None,
        authoritative: list[str] | None = None,
        generated_at: datetime | None = None,
        scanned_total: int = 0,
        excluded_total: int = 0,
    ):
        self.title = title
        self.classifier = classifier or CategoryClassifier([])
        self.authoritative = {a.strip() for a in (authoritative or []) if a.strip()}
        self.generated_at = generated_at or datetime.now(timezone.utc)

        # 只保留能归类的；归不了类的不进日报（用户要求「无关训练的就不要放进来」）
        relevant = [p for p in posts if p.relevance > 0]
        self.grouped, self.excluded = self.classifier.split(relevant)
        self.included: list[Post] = [p for items in self.grouped.values() for p in items]
        # 这两个是调用方给的全局口径（含被剔掉的），只用来写数据说明。
        # 生成器自己只能看到传进来的这批，算不出「本轮一共扫了多少」。
        self.scanned_total = scanned_total or len(relevant) + excluded_total
        self.excluded_total = excluded_total

    # ------------------------------------------------------------ 小工具

    @staticmethod
    def _metric(post: Post) -> tuple[int, str]:
        if post.platform == "wechat":
            return post.views, "阅读"
        return post.likes, "点赞"

    def _metric_text(self, post: Post) -> str:
        value, label = self._metric(post)
        return f"{label} {_fmt_count(value)}" if value else ""

    @staticmethod
    def _latest(posts: list[Post]) -> datetime | None:
        stamps = [p.publish_time for p in posts if p.publish_time]
        return max(stamps) if stamps else None

    # ------------------------------------------------------------ 样式

    @staticmethod
    def _configure_styles(doc) -> None:
        normal = doc.styles["Normal"]
        normal.font.name = EN_SERIF
        normal.font.size = Pt(SIZE_BODY)
        normal.font.color.rgb = _hex(Palette.body_text)
        rPr = normal.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        for attr in ("w:ascii", "w:hAnsi"):
            rFonts.set(qn(attr), EN_SERIF)
        rFonts.set(qn("w:eastAsia"), CN_SONG)
        normal.paragraph_format.line_spacing = LINE_SPACING

        # python-docx 默认模板的标题是蓝色 Calibri，必须改掉，否则标题色不统一
        for level, size in ((1, SIZE_H1), (2, SIZE_H2)):
            style = doc.styles[f"Heading {level}"]
            style.font.name = EN_SERIF
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = _hex(Palette.heading_text)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.line_spacing = LINE_SPACING
            sPr = style.element.get_or_add_rPr()
            sFonts = sPr.find(qn("w:rFonts"))
            if sFonts is None:
                sFonts = OxmlElement("w:rFonts")
                sPr.insert(0, sFonts)
            for attr in ("w:ascii", "w:hAnsi"):
                sFonts.set(qn(attr), EN_SERIF)
            sFonts.set(qn("w:eastAsia"), CN_HEI)
            _set_outline_level(style, level - 1)

    # ------------------------------------------------------------ 题头

    def _build_masthead(self, doc) -> None:
        """第一页顶部的标题区块（替代整页封面）。

        深绿底 + 主标题 + 一行统计，占几行的高度就够，
        不占一整页 —— 这是和「X 情报日报」对齐的地方。
        """
        table = doc.add_table(rows=1, cols=1)
        _no_borders(table)
        _cell_margins(table, top=180, bottom=180, left=200, right=200)

        row = table.rows[0]
        _row_cant_split(row)
        cell = row.cells[0]
        _shade(cell._tc, Palette.brand_bg)
        cell._tc.remove(cell.paragraphs[0]._p)

        # 主标题
        para = cell.add_paragraph()
        _set_spacing(para, after=60, line=520, rule=WD_LINE_SPACING.AT_LEAST)
        run = para.add_run(self.title)
        _set_run(run, cn=CN_HEI, en=EN_SANS, size=SIZE_DOC_TITLE, bold=True,
                 color=Palette.brand_text)

        # 日期 + 分类条数
        local = self.generated_at.astimezone(timezone(timedelta(hours=8)))
        parts = [local.strftime("%Y-%m-%d　%H:%M")]
        for key, items in self.grouped.items():
            name = next(
                (c.name for c in self.classifier.categories if c.key == key), key
            )
            parts.append(f"{name} {len(items)} 条")
        parts.append(f"合计 {len(self.included)} 条")

        para = cell.add_paragraph()
        _set_spacing(para, after=0)
        run = para.add_run("　·　".join(parts))
        _set_run(run, cn=CN_YAHEI, en=EN_SANS, size=SIZE_META + 1,
                 color=Palette.brand_sub)

        # 区块下方一条强调色横线，把题头和正文分开（用段落边框，不用字符画线）
        rule = doc.add_paragraph()
        _set_spacing(rule, before=0, after=200)
        _para_border(rule, edge="bottom", size=12, color=Palette.brand_accent, space=1)
        _set_table_widths(table, [TEXT_WIDTH_CM])

    # ------------------------------------------------------------ 正文构件

    @staticmethod
    def _heading(doc, text: str, level: int = 1):
        para = doc.add_paragraph(style=f"Heading {level}")
        run = para.add_run(text)
        size = SIZE_H1 if level == 1 else SIZE_H2
        _set_run(run, cn=CN_HEI, en=EN_SERIF, size=size, bold=True,
                 color=Palette.heading_text)
        return para

    @staticmethod
    def _body(doc, text: str, *, indent=True, color=Palette.body_text,
              size=SIZE_BODY, after=140):
        para = doc.add_paragraph()
        _set_spacing(para, after=after)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            para.paragraph_format.first_line_indent = BODY_INDENT
        run = para.add_run(text)
        _set_run(run, size=size, color=color)
        return para

    def _item(self, doc, index: int, post: Post, *, last: bool = False) -> None:
        """一条内容：标题 → 元信息 → 摘要 → 链接。

        条目之间用一条极浅的下边框分隔，长列表读起来不容易串行。
        """
        # 标题行
        para = doc.add_paragraph()
        _set_spacing(para, before=160, after=30)
        _keep_with_next(para)
        run = para.add_run(f"{index}. {_clean(post.title, 80) or '(无标题)'}")
        _set_run(run, cn=CN_HEI, en=EN_SERIF, size=SIZE_ITEM_TITLE, bold=True)

        # 元信息
        bits = [f"@{_clean(post.author)}" if post.author else "", platform_label(post.platform)]
        if post.author in self.authoritative:
            bits.append("权威号")
        bits += [_fmt_date(post.publish_time), self._metric_text(post)]
        para = doc.add_paragraph()
        _set_spacing(para, after=30)
        _keep_with_next(para)
        run = para.add_run("　·　".join(b for b in bits if b))
        _set_run(run, size=SIZE_META, color=Palette.meta_text)

        # 摘要
        summary = _clean(post.content, 150)
        if summary and summary != _clean(post.title):
            para = doc.add_paragraph()
            _set_spacing(para, after=30)
            para.paragraph_format.left_indent = BODY_INDENT
            run = para.add_run(summary)
            _set_run(run, size=SIZE_SUMMARY, color=Palette.summary_text)

        # 链接 + 该条目的下分隔线（最后一条不画，免得和下一节标题挤在一起）
        para = doc.add_paragraph()
        _set_spacing(para, after=0 if last else 120)
        if post.url:
            _add_hyperlink(para, post.url, "打开原文", size=SIZE_META)
        if not last:
            _para_border(para, edge="bottom", size=4, color=Palette.rule_light, space=6)

    def _overview_table(self, doc) -> None:
        """分类概览表。带最新发布时间，让「新不新」一眼可见。"""
        headers = ["分类", "条数", "最新发布", "来源账号数"]
        rows = []
        for key, items in self.grouped.items():
            name = next(
                (c.name for c in self.classifier.categories if c.key == key), key
            )
            latest = self._latest(items)
            rows.append([
                name,
                str(len(items)),
                _fmt_date(latest),
                str(len({p.author for p in items if p.author})),
            ])
        self._make_table(doc, headers, rows, [4.5, 2.0, 3.5, 3.5], left_cols=(0,))
        caption = doc.add_paragraph()
        _set_spacing(caption, before=60, after=200)
        run = caption.add_run("表 1　本期分类概览")
        _set_run(run, size=SIZE_META, color=Palette.meta_text)

    def _make_table(self, doc, headers, rows, widths, left_cols=()):
        table = doc.add_table(rows=1, cols=len(headers))
        _cell_margins(table)

        header_row = table.rows[0]
        _repeat_header(header_row)
        _row_cant_split(header_row)
        for idx, text in enumerate(headers):
            cell = header_row.cells[idx]
            _shade(cell._tc, Palette.table_header_bg)
            para = cell.paragraphs[0]
            para.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if idx in left_cols
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            _set_spacing(para)
            run = para.add_run(text)
            _set_run(run, cn=CN_HEI, en=EN_SERIF, size=SIZE_CAPTION, bold=True,
                     color=Palette.table_header_text)

        for row_data in rows:
            row = table.add_row()
            _row_cant_split(row)
            for idx, text in enumerate(row_data):
                cell = row.cells[idx]
                para = cell.paragraphs[0]
                para.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if idx in left_cols
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                _set_spacing(para)
                run = para.add_run(_clean(text))
                _set_run(run, size=SIZE_CAPTION)

        _set_table_widths(table, widths)
        return table

    # ------------------------------------------------------------ 组装

    def build(self) -> Document:
        doc = Document()
        self._configure_styles(doc)

        section = doc.sections[0]
        section.page_width, section.page_height = PAGE_W, PAGE_H
        section.top_margin, section.bottom_margin = MARGIN_TB, MARGIN_TB
        section.left_margin, section.right_margin = MARGIN_L, MARGIN_R
        _footer_page_number(section)

        self._build_masthead(doc)

        if not self.included:
            self._body(
                doc,
                "本期没有可收录的内容。用于判断的筛选条件是：命中「训练与技巧」"
                "或「队伍与赛事动态」两个分类之一；媒体评论、行业分析、装备种草"
                "这类内容按设定不会被收录。",
                indent=False,
            )
            return doc

        self._overview_table(doc)

        numbers = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7}
        seq = 1
        for key, items in self.grouped.items():
            category = next(
                (c for c in self.classifier.categories if c.key == key), None
            )
            name = category.name if category else key
            self._heading(doc, f"{list(numbers)[seq - 1]}、{name}（{len(items)} 条）")
            if category and category.intro:
                self._body(doc, category.intro, after=100)

            for i, post in enumerate(items, 1):
                self._item(doc, i, post, last=(i == len(items)))
            seq += 1

        # 数据说明
        self._heading(doc, f"{list(numbers)[seq - 1]}、本期数据说明")
        note = f"本轮扫描到 {self.scanned_total} 条飞盘相关内容，收录 {len(self.included)} 条。"
        if self.excluded_total:
            note += (
                f"另有 {self.excluded_total} 条属于媒体报道、行业分析、装备种草或"
                "蹭热点的内容，按设定未收录——只保留「训练与技巧」和"
                "「队伍与赛事动态」两类。"
            )
        self._body(doc, note)

        for text in (
            "采集范围：微信公众号（搜狗微信搜索）、小红书、抖音三个平台的公开内容。",
            "排序口径：每个分类内按发布时间从新到旧排列；跨分类按「训练与技巧 → "
            "队伍与赛事动态」的优先级排列。",
            "时效提醒：搜狗微信搜索按相关性而非时间返回结果，其索引对飞盘这类"
            "小众主题更新较慢，因此公众号部分可能出现往期文章，请以每条标注的"
            "发布日期为准。要稳定获取最新推送，建议改用公众号后台模式"
            "（config/sources.yaml 里把 wechat.mode 改成 mp）或直接依赖"
            "小红书 / 抖音——这两个平台支持按时间排序。",
            "版权声明：所有内容版权归原作者所有，本日报仅供内部选题与舆情参考，"
            "请勿二次分发或用于商业用途。",
        ):
            self._body(doc, text)

        return doc


def build_docx_daily(
    posts: list[Post],
    path: Path,
    *,
    title: str = "飞盘日报",
    classifier: CategoryClassifier | None = None,
    authoritative: list[str] | None = None,
    generated_at: datetime | None = None,
    scanned_total: int = 0,
    excluded_total: int = 0,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = DocxDailyReport(
        posts,
        title=title,
        classifier=classifier,
        authoritative=authoritative,
        generated_at=generated_at,
        scanned_total=scanned_total,
        excluded_total=excluded_total,
    ).build()
    doc.save(str(path))
    return path
